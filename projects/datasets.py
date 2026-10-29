"""Bounded CSV/XLSX import. Files and formulas are never executed or retained."""
import csv
import io
import math
import zipfile
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.core.exceptions import ValidationError
from django.db import transaction

from .models import Dataset, DatasetRow

MAX_BYTES = 2 * 1024 * 1024
MAX_ROWS = 1000
MAX_COLUMNS = 30
MAX_CELL_LENGTH = 2000


def normalize_value(raw, definition):
    """Shared scalar validation for uploaded and manually entered dataset values."""
    value = raw.strip()
    key = definition['key']
    if len(value) > MAX_CELL_LENGTH:
        raise ValidationError('Each cell must contain at most 2,000 characters.')
    error = None
    if not value and definition['required']:
        error = f'{key}: required value is missing.'
    elif value and definition['data_type'] == 'number':
        try:
            numeric = Decimal(value)
            if not numeric.is_finite() or len(numeric.as_tuple().digits) > 30 or abs(numeric.as_tuple().exponent) > 100:
                raise InvalidOperation
            value = str(numeric)
        except InvalidOperation:
            error = f'{key}: enter a finite number with at most 30 digits; use a decimal point, without thousands separators.'
    return value, error


def read_rows(upload):
    if upload.size > MAX_BYTES:
        raise ValidationError('File is too large. Maximum size is 2 MB.')
    content = upload.read(MAX_BYTES + 1)
    if len(content) > MAX_BYTES:
        raise ValidationError('File is too large. Maximum size is 2 MB.')
    extension = Path(upload.name).suffix.lower()
    if extension == '.csv':
        try:
            text = content.decode('utf-8-sig')
        except UnicodeDecodeError:
            raise ValidationError('CSV must use UTF-8 encoding. Choose CSV UTF-8 when exporting from Excel.')
        try:
            for row in csv.reader(io.StringIO(text, newline=''), strict=True):
                yield row
        except csv.Error:
            raise ValidationError('CSV could not be read. Use comma-separated columns and correctly quoted values.')
    elif extension == '.xlsx':
        from openpyxl import load_workbook
        book = None
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if len(archive.infolist()) > 1000 or sum(item.file_size for item in archive.infolist()) > 20 * 1024 * 1024:
                    raise ValidationError('Workbook is too large after decompression.')
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            if len(book.worksheets) != 1:
                raise ValidationError('Use a workbook with exactly one worksheet.')
            sheet = book.worksheets[0]
            # Ignore potentially incorrect worksheet dimension metadata; enforce limits ourselves.
            sheet.reset_dimensions()
            header_width = None
            for cells in sheet.iter_rows():
                if len(cells) > MAX_COLUMNS:
                    raise ValidationError('Maximum 30 columns per file.')
                row = []
                for cell in cells:
                    if cell.data_type in {'f', 'e'}:
                        raise ValidationError('Excel formulas and error cells are unsupported. Paste values before uploading.')
                    value = cell.value
                    if isinstance(value, (datetime, date, time)):
                        value = value.isoformat()
                    elif isinstance(value, bool):
                        value = 'true' if value else 'false'
                    elif isinstance(value, float) and not math.isfinite(value):
                        raise ValidationError('Excel contains a non-finite number.')
                    row.append('' if value is None else str(value))
                if header_width is None:
                    header_width = len(row)
                elif len(row) < header_width:
                    row.extend([''] * (header_width - len(row)))
                yield row
        except ValidationError:
            raise
        except Exception:
            raise ValidationError('Workbook could not be read. Upload a valid, unencrypted .xlsx file.')
        finally:
            if book:
                book.close()
    else:
        raise ValidationError('Upload a .csv or .xlsx file.')


def parse_dataset(upload, schema):
    if not schema:
        raise ValidationError('Define project inputs before uploading a dataset.')
    if len(schema) > MAX_COLUMNS:
        raise ValidationError('This MVP supports at most 30 project inputs per dataset.')
    rows = iter(read_rows(upload))
    try:
        headers = [value.strip() for value in next(rows)]
    except StopIteration:
        raise ValidationError('The file is empty. Include column headers and at least one data row.')
    if not headers or len(headers) > MAX_COLUMNS or any(not value for value in headers):
        raise ValidationError('Provide 1–30 non-empty column headers in the first row.')
    if len(set(headers)) != len(headers):
        raise ValidationError('Duplicate column headers are not allowed.')
    keys = [item['key'] for item in schema]
    if set(headers) != set(keys):
        raise ValidationError('Column headers must exactly match the project input keys: ' + ', '.join(keys))
    definitions = {item['key']: item for item in schema}
    records = []
    for number, row in enumerate(rows, start=2):
        if number > MAX_ROWS + 1:
            raise ValidationError('Maximum 1,000 data rows per file, including blank rows.')
        if not any(value.strip() for value in row):
            continue
        if len(row) != len(headers):
            raise ValidationError(f'Row {number} has a different number of columns from the header.')
        values, errors = {}, []
        for key, raw in zip(headers, row):
            try:
                value, error = normalize_value(raw, definitions[key])
            except ValidationError:
                raise ValidationError(f'Row {number}: each cell must contain at most 2,000 characters.') from None
            if error:
                errors.append(error)
            values[key] = value
        records.append({'row_number':number, 'values':values, 'errors':errors})
    if not records:
        raise ValidationError('Include at least one non-empty data row.')
    return headers, records


@transaction.atomic
def save_preview(project, name, source_name, schema, headers, records):
    dataset = Dataset.objects.create(project=project, name=name, source_name=source_name,
        schema_snapshot=schema, headers=headers, row_count=len(records), invalid_row_count=sum(bool(row['errors']) for row in records))
    DatasetRow.objects.bulk_create([DatasetRow(dataset=dataset, **record) for record in records])
    return dataset
